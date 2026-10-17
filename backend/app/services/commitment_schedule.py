"""Business commitment schedule — every loan a person or entity holds.

One schedule per holder (a contact, or an entity: company / trust / …), with
two sections in the same columns:

  ours  — loans we arranged, under the tenant's own name. A line is created
          from each settled deal the first time the schedule is drawn
          (`sync_our_loans`), prefilled from the deal and its lender pricing,
          and the desk's to correct from then on.
  other — commitments held elsewhere, keyed in by the desk. Who arranged them
          is never printed: the schedule does not distinguish our competitors
          from the lenders' own direct business.

An entity's schedule is its own loans only. A director's personal loans are on
the director's schedule; a company applicant's deal is on the company's.

Lines are LendingHistoryEntry rows (see that model). Columns are chosen per
download; the PPSR columns are off unless switched on.
"""
from __future__ import annotations

import calendar
import io
import json
from datetime import date, datetime
from typing import Iterable, Optional

from fastapi import HTTPException
from sqlalchemy.orm import Session

from app.models.contact import Contact, Organization
from app.models.lender import Lender
from app.models.lending_history_entry import LendingHistoryEntry, RepaymentFrequency
from app.models.loan_application import APPLICANT_TYPE_COMPANY, ApplicationStatus, LoanApplication
from app.models.settled_deal_snapshot import SettledDealParty, SettledDealSnapshot
from app.models.tenant import Tenant
from app.services.lender_pricing import latest_for_application
from app.services.loan_category import application_asset_details

# key, label, kind, shown by default. Kind drives formatting in both exports.
COLUMNS: list[dict] = [
    {"key": "sr", "label": "Sr no.", "kind": "int", "default": True},
    {"key": "settled_date", "label": "Settled date", "kind": "date", "default": True},
    {"key": "lender_name", "label": "Lender", "kind": "text", "default": True},
    {"key": "amount", "label": "Amount financed", "kind": "money", "default": True},
    {"key": "monthly_repayment", "label": "Monthly repayment", "kind": "money", "default": True},
    {"key": "term_months", "label": "Loan term", "kind": "term", "default": True},
    {"key": "balloon", "label": "Balloon", "kind": "money", "default": True},
    {"key": "interest_rate", "label": "Rate", "kind": "percent", "default": True},
    {"key": "asset_description", "label": "Asset details", "kind": "multiline", "default": True},
    {"key": "vin", "label": "VIN", "kind": "multiline", "default": True},
    {"key": "end_date", "label": "End date", "kind": "date", "default": True},
    {"key": "ppsr_charge_number", "label": "PPSR charge no.", "kind": "text", "default": False},
    {"key": "ppsr_change_number", "label": "PPSR change no.", "kind": "text", "default": False},
    {"key": "ppsr_organisation_identifier", "label": "Organisation identifier", "kind": "text", "default": False},
    {"key": "asset_price_ex_gst", "label": "Asset price ex GST", "kind": "money", "default": False},
    {"key": "gst_claimed", "label": "GST claimed", "kind": "money", "default": False},
    {"key": "tax_invoice_date", "label": "Tax invoice date", "kind": "date", "default": False},
]
COLUMN_KEYS = [c["key"] for c in COLUMNS]
_TOTALLED = ("amount", "monthly_repayment", "balloon")

# Fields the desk may edit on a line, ours or other.
EDITABLE_FIELDS = (
    "lender_name", "amount", "balloon", "repayment_amount", "repayment_frequency", "start_date",
    "interest_rate", "term_months", "end_date", "asset_description", "vin",
    "ppsr_charge_number", "ppsr_change_number", "ppsr_organisation_identifier",
    "asset_price_ex_gst", "gst_claimed", "tax_invoice_date", "notes", "other_broker_name",
)


# ── Holders ──────────────────────────────────────────────────────────

def get_holder(db: Session, holder_type: str, holder_id: str, tenant_id: str):
    if holder_type == "contact":
        holder = db.query(Contact).filter(Contact.id == holder_id, Contact.tenant_id == tenant_id).first()
    elif holder_type == "organization":
        holder = db.query(Organization).filter(Organization.id == holder_id, Organization.tenant_id == tenant_id).first()
    else:
        raise HTTPException(status_code=404, detail="Not found")
    if not holder:
        raise HTTPException(status_code=404, detail="Not found")
    return holder


def holder_name(holder) -> str:
    if isinstance(holder, Contact):
        return f"{holder.first_name} {holder.last_name}".strip()
    return holder.name


def _holder_filter(holder):
    if isinstance(holder, Contact):
        return LendingHistoryEntry.contact_id == holder.id
    return LendingHistoryEntry.organization_id == holder.id


# ── Our loans: created from settled deals ────────────────────────────

def _add_months(d: date, months: int) -> date:
    month = d.month - 1 + months
    year = d.year + month // 12
    month = month % 12 + 1
    return date(year, month, min(d.day, calendar.monthrange(year, month)[1]))


def _num(value) -> Optional[float]:
    try:
        n = float(value)
    except (TypeError, ValueError):
        return None
    return n


def _asset_text(details: dict) -> Optional[str]:
    if details.get("description"):
        return str(details["description"])
    parts = [str(details[k]) for k in ("condition", "year", "make", "model") if details.get(k)]
    if not parts and details.get("equipment_type"):
        parts = [str(details["equipment_type"])]
    return " ".join(parts) or None


def prefill_from_application(db: Session, app: LoanApplication) -> dict:
    """What a settled deal says about its loan. Only figures the desk actually
    recorded — the monthly repayment is taken only where the lender's figure was
    typed in, never re-derived, because a wrong number on a client's schedule is
    worse than a blank the desk fills in."""
    pricing = latest_for_application(db, app.id)
    params: dict = {}
    if pricing and pricing.input_parameters:
        try:
            params = json.loads(pricing.input_parameters) or {}
        except ValueError:
            params = {}
    lender = app.approval_lender_name or (pricing.lender.name if pricing and pricing.lender else None) or params.get("lender_name")
    term = params.get("term_months")
    term = int(term) if isinstance(term, (int, float)) and term > 0 else None
    balloon = _num(params.get("balloon_amount"))
    if balloon is None and params.get("balloon_on_total_price") and _num(params.get("balloon_percent")):
        balloon = round((_num(params.get("asset_price")) or 0) * _num(params["balloon_percent"]) / 100, 2)
    rate = _num(params.get("interest_rate"))
    monthly = _num(params.get("monthly_repayment_override_balloon") if balloon else params.get("monthly_repayment_override"))
    settled = app.settled_at.date() if app.settled_at else None
    asset = application_asset_details(app)
    return {
        "lender_name": lender or "—",
        "amount": float(app.amount),
        "balloon": balloon if balloon is not None else (0.0 if params else None),
        "repayment_amount": monthly if monthly and monthly > 0 else None,
        "repayment_frequency": RepaymentFrequency.monthly if monthly and monthly > 0 else None,
        "start_date": settled,
        "interest_rate": rate if rate and rate > 0 else None,
        "term_months": term,
        "end_date": _add_months(settled, term) if settled and term else None,
        "asset_description": _asset_text(asset) or params.get("asset_description") or None,
        "vin": asset.get("vin") or None,
        "asset_price_ex_gst": round(float(pricing.asset_price) / 1.1, 2) if pricing and pricing.asset_price else None,
    }


def _settled_deals(db: Session, holder, tenant_id: str) -> list[LoanApplication]:
    q = db.query(LoanApplication).filter(
        LoanApplication.tenant_id == tenant_id,
        LoanApplication.deleted_at.is_(None),
        LoanApplication.status == ApplicationStatus.settled,
    )
    if isinstance(holder, Contact):
        # The person borrowed it themselves — a company applicant's deal is the
        # company's, even though the person is its deal contact.
        q = q.filter(LoanApplication.contact_id == holder.id, LoanApplication.applicant_type != APPLICANT_TYPE_COMPANY)
    else:
        q = q.filter(LoanApplication.business_organization_id == holder.id, LoanApplication.applicant_type == APPLICANT_TYPE_COMPANY)
    return q.all()


def sync_our_loans(db: Session, holder, tenant_id: str) -> int:
    """Create a line for each settled deal of the holder that has none yet.
    Existing lines are never rewritten — the desk may have corrected them."""
    have = {
        app_id for (app_id,) in db.query(LendingHistoryEntry.application_id).filter(
            LendingHistoryEntry.tenant_id == tenant_id, LendingHistoryEntry.application_id.isnot(None),
        )
    }
    holder_cols = {"contact_id": holder.id} if isinstance(holder, Contact) else {"organization_id": holder.id}
    created = 0
    for app in _settled_deals(db, holder, tenant_id):
        if app.id in have:
            continue
        db.add(LendingHistoryEntry(tenant_id=tenant_id, application_id=app.id, **holder_cols, **prefill_from_application(db, app)))
        have.add(app.id)
        created += 1
    # Settled deals since archived and purged survive as snapshots naming the
    # person as main applicant; entities aren't on snapshots.
    if isinstance(holder, Contact):
        rows = (
            db.query(SettledDealSnapshot, Lender.name)
            .join(SettledDealParty, (SettledDealParty.application_id == SettledDealSnapshot.application_id)
                  & (SettledDealParty.tenant_id == SettledDealSnapshot.tenant_id))
            .outerjoin(Lender, (Lender.id == SettledDealSnapshot.lender_id) & (Lender.tenant_id == tenant_id))
            .filter(SettledDealSnapshot.tenant_id == tenant_id, SettledDealParty.contact_id == holder.id,
                    SettledDealParty.role == "Main applicant")
            .all()
        )
        for snap, lender in rows:
            if snap.application_id in have:
                continue
            db.add(LendingHistoryEntry(
                tenant_id=tenant_id, application_id=snap.application_id, contact_id=holder.id,
                lender_name=lender or "—", amount=float(snap.amount),
                start_date=snap.settled_at.date() if snap.settled_at else None,
            ))
            have.add(snap.application_id)
            created += 1
    if created:
        db.flush()
    return created


# ── The schedule ─────────────────────────────────────────────────────

def monthly_repayment(entry: LendingHistoryEntry) -> Optional[float]:
    if entry.repayment_amount is None:
        return None
    freq = entry.repayment_frequency or RepaymentFrequency.monthly
    if freq == RepaymentFrequency.weekly:
        return round(entry.repayment_amount * 52 / 12, 2)
    if freq == RepaymentFrequency.fortnightly:
        return round(entry.repayment_amount * 26 / 12, 2)
    return entry.repayment_amount


def _row(entry: LendingHistoryEntry, sr: int) -> dict:
    missing = [label for label, value in (
        ("monthly repayment", entry.repayment_amount), ("term", entry.term_months), ("rate", entry.interest_rate),
    ) if value is None]
    return {
        "id": entry.id,
        "application_id": entry.application_id,
        "sr": sr,
        "settled_date": entry.start_date,
        "lender_name": entry.lender_name,
        "amount": entry.amount,
        "monthly_repayment": monthly_repayment(entry),
        "term_months": entry.term_months,
        "balloon": entry.balloon,
        "interest_rate": entry.interest_rate,
        "asset_description": entry.asset_description,
        "vin": entry.vin,
        "end_date": entry.end_date,
        "ppsr_charge_number": entry.ppsr_charge_number,
        "ppsr_change_number": entry.ppsr_change_number,
        "ppsr_organisation_identifier": entry.ppsr_organisation_identifier,
        "asset_price_ex_gst": entry.asset_price_ex_gst,
        "gst_claimed": entry.gst_claimed,
        "tax_invoice_date": entry.tax_invoice_date,
        # Editing only — never exported.
        "edit": {
            "repayment_amount": entry.repayment_amount,
            "repayment_frequency": entry.repayment_frequency.value if entry.repayment_frequency else None,
            "notes": entry.notes,
            "other_broker_name": entry.other_broker_name,
        },
        "missing": missing,
    }


def _totals(rows: Iterable[dict]) -> dict:
    rows = list(rows)
    totals = {k: round(sum(r[k] or 0 for r in rows), 2) for k in _TOTALLED}
    totals["annual_repayment"] = round(totals["monthly_repayment"] * 12, 2)
    return totals


def build(db: Session, holder, tenant_id: str) -> dict:
    sync_our_loans(db, holder, tenant_id)
    db.commit()
    entries = (
        db.query(LendingHistoryEntry)
        .filter(LendingHistoryEntry.tenant_id == tenant_id, _holder_filter(holder))
        .order_by(LendingHistoryEntry.start_date.asc().nullslast(), LendingHistoryEntry.created_at.asc())
        .all()
    )
    tenant = db.query(Tenant).filter(Tenant.id == tenant_id).first()
    arranger = tenant.name if tenant else "Us"
    ours = [e for e in entries if e.application_id]
    other = [e for e in entries if not e.application_id]
    sections = []
    for key, title, group in (("ours", f"Arranged by {arranger}", ours), ("other", "Other commitments", other)):
        rows = [_row(e, i + 1) for i, e in enumerate(group)]
        sections.append({"key": key, "title": title, "rows": rows, "totals": _totals(rows)})
    all_rows = [r for s in sections for r in s["rows"]]
    return {
        "holder": {
            "type": "contact" if isinstance(holder, Contact) else "organization",
            "id": holder.id,
            "name": holder_name(holder),
            "abn": getattr(holder, "abn", None),
        },
        "arranger": arranger,
        "generated_at": datetime.now().date(),
        "columns": COLUMNS,
        "sections": sections,
        "totals": _totals(all_rows),
    }


def parse_columns(value: Optional[str]) -> list[str]:
    if not value:
        return [c["key"] for c in COLUMNS if c["default"]]
    wanted = {k.strip() for k in value.split(",")}
    keys = [k for k in COLUMN_KEYS if k in wanted]
    if not keys:
        raise HTTPException(status_code=400, detail="Choose at least one column")
    return keys


# ── Excel ────────────────────────────────────────────────────────────

def to_xlsx(schedule: dict, columns: list[str]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    spec = {c["key"]: c for c in COLUMNS}
    cols = [spec[k] for k in columns]
    wb = Workbook()
    ws = wb.active
    ws.title = "Commitment schedule"
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="1F2A44")
    band_fill = PatternFill("solid", fgColor="EEF1F6")
    thin = Side(style="thin", color="C9CED8")
    box = Border(top=thin, bottom=thin, left=thin, right=thin)
    fmt = {"money": '"$"#,##0.00', "percent": "0.00%", "date": "DD/MM/YYYY", "int": "0"}

    ws.cell(row=1, column=1, value="Business Commitment Schedule").font = Font(bold=True, size=14)
    holder = schedule["holder"]
    ws.cell(row=2, column=1, value=holder["name"] + (f"  ·  ABN {holder['abn']}" if holder.get("abn") else "")).font = bold
    ws.cell(row=3, column=1, value=f"Prepared {schedule['generated_at'].strftime('%d/%m/%Y')}")
    r = 5

    # Totals labels sit in the first column that isn't itself a total.
    label_col = next((i for i, c in enumerate(cols, start=1) if c["key"] not in _TOTALLED), 1)

    def write_totals(label: str, totals: dict, row: int, strong: bool = False) -> None:
        ws.cell(row=row, column=label_col, value=label).font = bold
        for i, c in enumerate(cols, start=1):
            if c["key"] in _TOTALLED:
                cell = ws.cell(row=row, column=i, value=totals[c["key"]])
                cell.number_format = fmt["money"]
                cell.font = bold
            if strong:
                ws.cell(row=row, column=i).fill = band_fill

    for section in schedule["sections"]:
        ws.cell(row=r, column=1, value=section["title"]).font = Font(bold=True, size=12)
        r += 1
        for i, c in enumerate(cols, start=1):
            cell = ws.cell(row=r, column=i, value=c["label"])
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = head_fill
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            cell.border = box
        r += 1
        if not section["rows"]:
            ws.cell(row=r, column=1, value="None recorded.").font = Font(italic=True, color="808080")
            r += 2
            continue
        for row in section["rows"]:
            for i, c in enumerate(cols, start=1):
                value = row[c["key"]]
                if c["kind"] == "percent" and value is not None:
                    value = value / 100
                elif c["kind"] == "term" and value is not None:
                    value = f"{value} months"
                cell = ws.cell(row=r, column=i, value=value)
                cell.border = box
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if c["kind"] in fmt:
                    cell.number_format = fmt[c["kind"]]
            r += 1
        write_totals("Subtotal", section["totals"], r)
        r += 2

    write_totals("TOTAL", schedule["totals"], r, strong=True)
    ws.cell(row=r + 1, column=label_col, value="Annual repayments").font = bold
    annual_col = next((i for i, c in enumerate(cols, start=1) if c["key"] == "monthly_repayment"), 2)
    cell = ws.cell(row=r + 1, column=annual_col, value=schedule["totals"]["annual_repayment"])
    cell.number_format = fmt["money"]
    cell.font = bold

    widths = {"int": 7, "date": 12, "money": 15, "percent": 8, "term": 11, "text": 22, "multiline": 38}
    for i, c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths[c["kind"]]
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
